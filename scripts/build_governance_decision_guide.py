#!/usr/bin/env python3
"""Build the stakeholder-facing repository governance decision guide DOCX.

The Markdown file remains the versioned source of truth.  This builder applies
the compact-reference-guide document preset and fails if the 41 decision cards
are not present exactly once.
"""

from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path
import posixpath
import re
from urllib.parse import quote, urlparse

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


VERSION = "0.1.0-draft"
TITLE = "Repository Governance Decision Guide"
SUBTITLE = "Forty-one decisions that make repository-health assessment repeatable, automatable, and accountable"
OWNER = "Enterprise Architecture"
EXPECTED_DECISIONS = tuple(range(1, 42))
REPOSITORY_BLOB_BASE = "https://github.com/TedTschopp/github-repository-health/blob/main/"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "203040"
MUTED = "5D6B78"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF4D6"
PALE_GRAY = "F3F5F7"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_table_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_paragraph_border(paragraph, *, color: str = BLUE, size: int = 16, space: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    borders.append(left)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))


def set_explicit_list_number(paragraph, value: int) -> None:
    """Apply a real Word list number while preserving the Markdown number."""
    numbering = paragraph.part.numbering_part.element
    existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    # The stock python-docx template's abstract numbering 7 is List Number.
    abstract.set(qn("w:val"), "7")
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), str(value))
    override.append(start)
    num.append(override)
    numbering.append(num)

    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    number_id = OxmlElement("w:numId")
    number_id.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, number_id))


def add_hyperlink(paragraph, text: str, url: str, *, color: str = BLUE, underline: bool = True) -> None:
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    run_color = OxmlElement("w:color")
    run_color.set(qn("w:val"), color)
    run_props.append(run_color)
    if underline:
        underline_node = OxmlElement("w:u")
        underline_node.set(qn("w:val"), "single")
        run_props.append(underline_node)
    run.append(run_props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_TOKEN = re.compile(
    r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://[^\s)>]+)"
)


def add_inline(paragraph, text: str) -> None:
    """Render a deliberately small, safe Markdown inline subset."""
    cursor = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        elif token.startswith("["):
            link_match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_match:
                label, target = link_match.groups()
                parsed = urlparse(target)
                if parsed.scheme in {"http", "https"}:
                    add_hyperlink(paragraph, label, target)
                elif target.startswith("#"):
                    paragraph.add_run(label)
                else:
                    path, separator, fragment = target.partition("#")
                    normalized = posixpath.normpath(posixpath.join("docs", path))
                    url = REPOSITORY_BLOB_BASE + quote(normalized, safe="/")
                    if separator:
                        url += "#" + quote(fragment, safe="-_.")
                    add_hyperlink(paragraph, label, url)
        else:
            add_hyperlink(paragraph, token, token)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Title", 27, DARK_BLUE, 0, 12),
        ("Subtitle", 14, MUTED, 0, 16),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
        ("Heading 4", 11, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Decision Recommendation" not in styles:
        style = styles.add_style("Decision Recommendation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Decision Recommendation"]
    style.base_style = styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.bold = False
    style.paragraph_format.left_indent = Inches(0.14)
    style.paragraph_format.right_indent = Inches(0.08)
    style.paragraph_format.space_before = Pt(5)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.keep_together = True

    if "Decision Label" not in styles:
        style = styles.add_style("Decision Label", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Decision Label"]
    style.base_style = styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    style.paragraph_format.space_before = Pt(5)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.keep_with_next = True


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.text = f"{TITLE}  |  {VERSION}"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.add_run("Enterprise Architecture  •  Provisional decision-support artifact")
    p.add_run("\t")
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(6.25), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)
    add_page_number(p)
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    first_footer = section.first_page_footer
    first_p = first_footer.paragraphs[0]
    first_p.text = "0.1.0-draft  •  Workshop use  •  Not a compliance attestation"
    first_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in first_p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    document.core_properties.title = TITLE
    document.core_properties.subject = "Repository governance workshop decision guide"
    document.core_properties.author = OWNER
    document.core_properties.keywords = "Git, repository health, governance, GitHub, decision guide"
    document.core_properties.comments = "Generated from docs/repository-governance-decision-guide.md"


def add_cover(document: Document) -> None:
    accent = document.add_paragraph()
    accent.paragraph_format.space_after = Pt(22)
    run = accent.add_run("GIT REPOSITORY HEALTH STANDARD")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    set_paragraph_border(accent, color=BLUE, size=22, space=8)

    title = document.add_paragraph(style="Title")
    title.add_run(TITLE)
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(SUBTITLE)

    note = document.add_paragraph(style="Decision Recommendation")
    note.paragraph_format.space_after = Pt(18)
    set_paragraph_border(note, color=BLUE, size=18, space=7)
    add_inline(
        note,
        "**Purpose.** Give accountable leaders a finite, explainable set of choices that lets evidence collection become repeatable without turning every assessment into a new interview.",
    )

    metadata = document.add_table(rows=6, cols=2)
    metadata.style = "Table Grid"
    metadata.alignment = WD_TABLE_ALIGNMENT.LEFT
    metadata.autofit = False
    values = (
        ("Version", VERSION),
        ("Status", "Provisional decision-support guide"),
        ("Accountable standard owner", OWNER),
        ("Prepared for", "Repository owners, engineering, security, operations, risk, and architecture"),
        ("Prepared", date.today().isoformat()),
    )
    metadata.rows[0].cells[0].paragraphs[0].add_run("Document field").bold = True
    metadata.rows[0].cells[1].paragraphs[0].add_run("Value").bold = True
    set_cell_width(metadata.rows[0].cells[0], 2400)
    set_cell_width(metadata.rows[0].cells[1], 6960)
    set_repeat_table_header(metadata.rows[0])
    keep_table_row_together(metadata.rows[0])
    for cell in metadata.rows[0].cells:
        set_cell_shading(cell, PALE_BLUE)
        set_cell_margins(cell, top=90, bottom=90)
    for row, (label, value) in zip(metadata.rows[1:], values, strict=True):
        keep_table_row_together(row)
        set_cell_width(row.cells[0], 2400)
        set_cell_width(row.cells[1], 6960)
        set_cell_shading(row.cells[0], PALE_BLUE)
        for cell in row.cells:
            set_cell_margins(cell, top=90, bottom=90)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(value)

    caveat = document.add_paragraph()
    caveat.paragraph_format.space_before = Pt(18)
    caveat.paragraph_format.keep_together = True
    run = caveat.add_run("Use boundary")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    caveat.add_run(
        ": this guide helps governance bodies make and record policy choices. It does not itself approve a choice, configure GitHub, accept risk, or prove that a control operated."
    )

    document.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(cells)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def table_widths(column_count: int) -> list[int]:
    if column_count == 2:
        return [2500, 6860]
    if column_count == 3:
        return [900, 3300, 5160]
    if column_count == 4:
        return [700, 2600, 2200, 3860]
    base = 9360 // column_count
    widths = [base] * column_count
    widths[-1] += 9360 - sum(widths)
    return widths


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = table_widths(column_count)
    for row_index, source_row in enumerate(rows):
        target_row = table.rows[row_index]
        keep_table_row_together(target_row)
        target_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for column_index in range(column_count):
            cell = target_row.cells[column_index]
            set_cell_width(cell, widths[column_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_index == 0:
                set_cell_shading(cell, PALE_BLUE)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_GRAY)
            text = source_row[column_index] if column_index < len(source_row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, text)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    set_repeat_table_header(table.rows[0])


def add_decision_paragraph(document: Document, text: str) -> None:
    label_match = re.match(r"\*\*([^*]+)\*\*\s*(.*)$", text)
    label = label_match.group(1).rstrip(".:") if label_match else ""
    body = label_match.group(2).lstrip(" .:") if label_match else text
    if label.lower() == "recommended default":
        paragraph = document.add_paragraph(style="Decision Recommendation")
        set_paragraph_border(paragraph, color=BLUE, size=16, space=6)
        add_inline(paragraph, text)
        return
    paragraph = document.add_paragraph()
    if label_match:
        run = paragraph.add_run(f"{label}. ")
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        add_inline(paragraph, body)
    else:
        add_inline(paragraph, text)


def add_markdown(document: Document, source: str) -> None:
    lines = source.splitlines()
    # The cover owns the H1. Start at the first H2 to avoid duplicating source metadata.
    first_h2 = next((i for i, line in enumerate(lines) if line.startswith("## ")), 0)
    lines = lines[first_h2:]
    index = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            add_decision_paragraph(document, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer = []

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        heading_match = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            hashes, heading_text = heading_match.groups()
            level = len(hashes)
            paragraph = document.add_paragraph(style=f"Heading {level}")
            add_inline(paragraph, heading_text)
            if re.match(r"^\d{1,2}\.\s", heading_text):
                paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue
        if stripped.startswith("> "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="Decision Recommendation")
            set_paragraph_border(paragraph, color="D6A533", size=16, space=6)
            add_inline(paragraph, stripped[2:])
            index += 1
            continue
        bullet_match = re.match(r"^[-*]\s+(.+)$", stripped)
        number_match = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if bullet_match or number_match:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet" if bullet_match else "List Number")
            if bullet_match:
                add_inline(paragraph, bullet_match.group(1))
            else:
                set_explicit_list_number(paragraph, int(number_match.group(1)))
                add_inline(paragraph, number_match.group(2))
            index += 1
            continue
        if stripped in {"---", "***"}:
            flush_paragraph()
            separator = document.add_paragraph()
            separator.paragraph_format.space_before = Pt(3)
            separator.paragraph_format.space_after = Pt(5)
            p_pr = separator._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "C8D4E0")
            borders.append(bottom)
            p_pr.append(borders)
            index += 1
            continue
        paragraph_buffer.append(stripped)
        index += 1
    flush_paragraph()


def audit_source(source: str) -> None:
    decisions = [int(value) for value in re.findall(r"^##\s+(\d{1,2})\.\s+", source, flags=re.MULTILINE)]
    if tuple(decisions) != EXPECTED_DECISIONS:
        raise ValueError(
            "Decision headings must be numbered 1 through 41 exactly once and in order; "
            f"found {decisions}"
        )
    required_labels = (
        "Decision question",
        "What must be decided",
        "Why it matters",
        "Acceptable values",
        "Recommended default",
        "Decision authority",
        "Record in",
        "Evidence and review",
        "Source basis",
    )
    for decision in EXPECTED_DECISIONS:
        match = re.search(
            rf"^##\s+{decision}\.\s+.*?(?=^##\s+{decision + 1}\.\s+|^#\s|\Z)",
            source,
            flags=re.MULTILINE | re.DOTALL,
        )
        if not match:
            raise ValueError(f"Decision {decision} is missing")
        card = match.group(0)
        missing = [label for label in required_labels if label.lower() not in card.lower()]
        if missing:
            raise ValueError(f"Decision {decision} is missing required labels: {', '.join(missing)}")


def build(source_path: Path, output_path: Path) -> None:
    source = source_path.read_text(encoding="utf-8")
    audit_source(source)
    document = Document()
    configure_styles(document)
    configure_document(document)
    add_cover(document)
    add_markdown(document, source)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        type=Path,
        default=Path("docs/repository-governance-decision-guide.md"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("deliverables/Repository Governance Decision Guide 0.1.0-draft.docx"),
    )
    args = parser.parse_args()
    build(args.source, args.output)
    print(f"Built {args.output} from {args.source} with 41 validated decision cards.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
